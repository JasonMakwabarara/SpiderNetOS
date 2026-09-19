"""Finance model — golden values, live-formula verification, validation, endpoint."""

from __future__ import annotations

import base64
import io

import pytest

from intelligence.services.business_launch import DISCLAIMER, finance_model as fm
from intelligence.services.business_launch.formula_eval import FormulaError, Workbook, column_index, column_letter, evaluate_workbook

GROWING = {
    "currency": "GBP",
    "starting_cash": 10000,
    "setup_costs": 2000,
    "price_per_unit": 100,
    "units_month_one": 20,
    "monthly_growth_pct": 10,
    "cost_of_sale_pct": 40,
    "fixed_monthly_costs": 1500,
    "marketing_monthly": 300,
    "months": 12,
}

RUNS_OUT = {
    "currency": "ZAR",
    "starting_cash": 5000,
    "setup_costs": 1000,
    "price_per_unit": 50,
    "units_month_one": 10,
    "monthly_growth_pct": 0,
    "cost_of_sale_pct": 50,
    "fixed_monthly_costs": 1000,
    "months": 12,
}


def test_golden_summary_for_a_growing_business():
    result = fm.run_model(GROWING, include_workbook=False)
    s = result["summary"]

    # Hand-checked: revenue = 2000 × (1.1^12 − 1) / 0.1; margin 60%; EBITDA
    # = gross − 12 × (1500 + 300) − 2000 set-up; cash low after month 5.
    assert s["currency"] == "GBP"
    assert s["months"] == 12
    assert s["year1_revenue"] == 42768.57
    assert s["year1_gross_profit"] == 25661.14
    assert s["year1_gross_margin_pct"] == 60.0
    assert s["year1_ebitda"] == 2061.14
    assert s["month1_burn"] == 2600.0
    assert s["cash_low_point"] == 6326.12
    assert s["cash_low_month"] == 5
    assert s["runs_out"] is False
    assert s["runway_months"] == 12
    assert s["breakeven_month"] == 6
    assert s["closing_cash"] == 12061.14
    assert s["years"] == [{
        "year": 1, "months": 12, "revenue": 42768.57, "gross_profit": 25661.14,
        "gross_margin_pct": 60.0, "ebitda": 2061.14, "closing_cash": 12061.14,
    }]
    assert s["disclaimer"] == DISCLAIMER == "Not legal or financial advice."
    assert result["disclaimer"] == DISCLAIMER

    month1 = result["monthly"][0]
    assert month1["revenue"] == 2000.0
    assert month1["ebitda"] == -2600.0
    assert month1["closing_cash"] == 7400.0


def test_golden_summary_when_cash_runs_out():
    s = fm.run_model(RUNS_OUT, include_workbook=False)["summary"]

    # EBITDA is −750 every month (−1,750 in month 1 with set-up costs).
    assert s["year1_revenue"] == 6000.0
    assert s["year1_gross_margin_pct"] == 50.0
    assert s["year1_ebitda"] == -10000.0
    assert s["runs_out"] is True
    assert s["runway_months"] == 5          # month 6 closes at −500
    assert s["cash_low_point"] == -5000.0
    assert s["cash_low_month"] == 12
    assert s["breakeven_month"] is None


def test_multi_year_table_and_default_horizon():
    assumptions = dict(GROWING)
    assumptions.pop("months")
    s = fm.run_model(assumptions, include_workbook=False)["summary"]
    assert s["months"] == 36
    assert [y["year"] for y in s["years"]] == [1, 2, 3]
    assert s["years"][0]["revenue"] == 42768.57
    assert s["years"][2]["closing_cash"] == s["closing_cash"]


def test_workbook_formulas_are_live_and_reproduce_the_projection():
    a = fm.Assumptions.from_mapping(GROWING)
    spec = fm.build_workbook_spec(a)

    assert spec["order"] == ["Assumptions", "Revenue", "Costs", "P&L", "CashFlow", "Balance"]
    revenue = spec["sheets"]["Revenue"]
    assert revenue[0] == ["Month", "Sales", "Price", "Revenue"]
    assert revenue[1] == [1, "=Assumptions!$B$6*(1+Assumptions!$B$7)^(A2-1)", "=Assumptions!$B$5", "=B2*C2"]
    assert revenue[13] == ["Total", "=SUM(B2:B13)", "", "=SUM(D2:D13)"]
    assert spec["sheets"]["Costs"][1][4] == "=IF(A2=1,Assumptions!$B$4,0)"
    assert spec["sheets"]["CashFlow"][1][1] == "=Assumptions!$B$3"
    assert spec["sheets"]["CashFlow"][2][1] == "=D2"
    assert spec["sheets"]["CashFlow"][2][2] == "='P&L'!F3"
    assert spec["sheets"]["Balance"][2][4] == "=E2+'P&L'!F3"
    # Percentages live as fractions so the formulas stay plain arithmetic.
    assert spec["sheets"]["Assumptions"][6] == ["Monthly sales growth", 0.1, "Compounded month on month."]
    assert DISCLAIMER.rstrip(".") in spec["sheets"]["Assumptions"][-1][0]

    rows = fm.project(a)
    verification = fm.verify_workbook(spec, rows)
    assert verification["verified"] is True, verification["mismatches"]
    assert verification["cells_checked"] > 200

    book = Workbook(spec["sheets"])
    assert round(book.value("Revenue", "D14"), 2) == 42768.57          # Total revenue
    assert round(book.value("CashFlow", "B15"), 2) == 6326.12          # Lowest closing cash
    assert round(book.value("CashFlow", "D13"), 2) == 12061.14
    for row in range(2, 14):
        assert abs(book.value("Balance", f"G{row}")) < 1e-6              # balance sheet balances


def test_verification_catches_a_broken_formula():
    a = fm.Assumptions.from_mapping(GROWING)
    spec = fm.build_workbook_spec(a)
    spec["sheets"]["P&L"][3][5] = "=D4-E4+1"  # month 3 EBITDA off by one
    verification = fm.verify_workbook(spec, fm.project(a))
    assert verification["verified"] is False
    assert any(m["sheet"] == "P&L" and m["cell"] == "F4" for m in verification["mismatches"])


def test_run_model_reports_workbook_availability():
    result = fm.run_model(GROWING)
    assert result["formulas_verified"] is True
    workbook = result["workbook"]
    assert workbook["format"] == "xlsx"
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        assert workbook["available"] is False
        assert "openpyxl" in workbook["reason"]
    else:
        assert workbook["available"] is True
        assert workbook["bytes"] > 0 and len(workbook["sha256"]) == 64


def test_xlsx_keeps_formulas_when_openpyxl_is_installed():
    openpyxl = pytest.importorskip("openpyxl")
    a = fm.Assumptions.from_mapping(GROWING)
    data, reason = fm.workbook_bytes(fm.build_workbook_spec(a))
    assert reason is None and data
    book = openpyxl.load_workbook(io.BytesIO(data))  # formulas, not cached values
    assert book.sheetnames == ["Assumptions", "Revenue", "Costs", "P&L", "CashFlow", "Balance"]
    assert book["Revenue"]["D2"].value == "=B2*C2"
    assert book["CashFlow"]["C3"].value == "='P&L'!F3"
    assert book["Assumptions"]["B7"].number_format == "0.00%"


@pytest.mark.parametrize(
    ("patch", "field"),
    [
        ({"currency": "pounds"}, "currency"),
        ({"starting_cash": None}, "starting_cash"),
        ({"price_per_unit": "a lot"}, "price_per_unit"),
        ({"cost_of_sale_pct": -5}, "cost_of_sale_pct"),
        ({"months": 0}, "months"),
        ({"months": 2.5}, "months"),
        ({"units_month_one": float("inf")}, "units_month_one"),
    ],
)
def test_invalid_assumptions_are_rejected_with_field_errors(patch, field):
    with pytest.raises(fm.AssumptionError) as excinfo:
        fm.Assumptions.from_mapping({**GROWING, **patch})
    assert field in excinfo.value.errors


def test_numeric_strings_are_accepted_and_marketing_is_optional():
    data = {k: str(v) for k, v in GROWING.items() if k != "marketing_monthly"}
    a = fm.Assumptions.from_mapping(data)
    assert a.marketing_monthly == 0.0 and a.months == 12 and a.currency == "GBP"


def test_round2_is_half_up():
    assert fm.round2(2.675) == 2.68
    assert fm.round2(-1.005) == -1.01


def test_formula_evaluator_basics():
    assert column_index("A") == 1 and column_index("AA") == 27 and column_letter(28) == "AB"
    values = evaluate_workbook({
        "S": [
            [2, 3, "=A1^B1", "=-A1^2", "=IF(A1>1,\"big\",\"small\")"],
            ["=SUM(A1:B1)", "=MIN(A1:B1)", "=MAX(A1:C1)", "=ROUND(2.5,0)", "=A1&\"x\""],
        ],
        "Other sheet": [["='S'!C1/2"]],
    })
    assert values["S"][0][2:] == [8.0, 4.0, "big"]  # Excel: -2^2 = 4
    assert values["S"][1] == [5.0, 2.0, 8.0, 3.0, "2x"]
    assert values["Other sheet"][0][0] == 4.0
    with pytest.raises(FormulaError):
        evaluate_workbook({"S": [["=B1", "=A1"]]})  # circular


def test_finance_model_endpoint():
    fastapi = pytest.importorskip("fastapi")  # noqa: F841
    from fastapi.testclient import TestClient

    from intelligence.services.business_launch.app import app

    client = TestClient(app)
    ok = client.post("/finance/model", json={"assumptions": GROWING, "include_workbook": False})
    assert ok.status_code == 200
    body = ok.json()
    assert body["summary"]["cash_low_point"] == 6326.12
    assert body["formulas_verified"] is True
    assert body["disclaimer"] == DISCLAIMER

    bad = client.post("/finance/model", json={"assumptions": {**GROWING, "currency": ""}})
    assert bad.status_code == 422
    assert "currency" in bad.json()["detail"]["errors"]

    health = client.get("/health").json()
    assert health["status"] == "ok"
    assert health["jurisdictions"] == ["uk", "za", "zw"]


def test_internal_key_is_enforced_when_configured(monkeypatch):
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient

    from intelligence.services.business_launch.app import app

    monkeypatch.setenv("BACKEND_INTERNAL_KEY", "s3cret")
    client = TestClient(app)
    assert client.post("/finance/model", json={"assumptions": GROWING}).status_code == 401
    assert client.post(
        "/finance/model", json={"assumptions": GROWING, "include_workbook": False}, headers={"X-Internal-Key": "s3cret"}
    ).status_code == 200
    assert client.get("/health").status_code == 200


def test_b64_workbook_round_trips_when_available():
    result = fm.run_model(GROWING)
    if result["workbook"]["available"]:
        assert len(base64.b64decode(result["workbook"]["b64"])) == result["workbook"]["bytes"]
    else:
        assert "b64" not in result["workbook"]
