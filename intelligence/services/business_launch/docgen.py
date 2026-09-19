"""Document generation for the business-launch pack.

Jinja template → markdown (always) → .docx (python-docx, only when importable)
→ .pdf (weasyprint or pandoc, only when present). Optional renderers are
imported or located inside the functions that use them, never at import time,
so the service starts and tests run on a machine with none of them installed;
a format that cannot be produced is reported with a reason instead of failing.

Every rendered document starts with the "Not legal or financial advice." line —
``render_markdown`` adds it if a template ever forgets.
"""

from __future__ import annotations

import base64
import html
import importlib.util
import re
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path
from typing import Any, Mapping, Sequence

from . import DISCLAIMER, pack_dir

__all__ = [
    "DocgenError",
    "SUPPORTED_FORMATS",
    "capabilities",
    "format_money",
    "format_pct",
    "markdown_to_docx",
    "markdown_to_html",
    "markdown_to_pdf",
    "render",
    "render_markdown",
]

SUPPORTED_FORMATS = ("md", "docx", "pdf")
DEFAULT_TEMPLATE = "business-plan.md.j2"
TO_CONFIRM = "[to confirm]"
_TEMPLATE_NAME_RE = re.compile(r"^[a-z0-9][a-z0-9-]*\.md\.j2$")
_SYMBOLS = {"GBP": "£", "ZAR": "R", "USD": "$", "EUR": "€"}


class DocgenError(ValueError):
    """Bad template name, missing template or unsupported format."""


def _is_missing(value: Any) -> bool:
    from jinja2 import Undefined

    return value is None or value == "" or isinstance(value, Undefined)


def format_money(value: Any, currency: Any = "") -> str:
    """``42768.567`` + ``GBP`` → ``£42,768.57``; negatives as ``-£5,000.00``."""
    if _is_missing(value):
        return TO_CONFIRM
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    code = "" if _is_missing(currency) else str(currency).upper()
    symbol = _SYMBOLS.get(code)
    body = f"{abs(number):,.2f}"
    text = f"{symbol}{body}" if symbol else (f"{code} {body}" if code else body)
    return f"-{text}" if number < 0 else text


def format_pct(value: Any) -> str:
    if _is_missing(value):
        return "n/a"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    return f"{number:,.2f}".rstrip("0").rstrip(".") + "%"


def _template_dir(template_dir: str | Path | None) -> Path:
    return Path(template_dir) if template_dir else pack_dir() / "templates"


def render_markdown(
    context: Mapping[str, Any],
    template: str = DEFAULT_TEMPLATE,
    *,
    template_dir: str | Path | None = None,
) -> str:
    from jinja2 import ChainableUndefined, Environment, FileSystemLoader

    if not _TEMPLATE_NAME_RE.match(template):
        raise DocgenError(f"Invalid template name: {template!r}")
    folder = _template_dir(template_dir)
    if not (folder / template).is_file():
        raise DocgenError(f"Template not found: {template}")

    env = Environment(
        loader=FileSystemLoader(str(folder)),
        undefined=ChainableUndefined,
        autoescape=False,
        keep_trailing_newline=True,
    )
    env.filters["money"] = format_money
    env.filters["pct"] = format_pct

    data = dict(context)
    data.setdefault("disclaimer", None)
    data.setdefault("generated_on", date.today().isoformat())
    text = env.get_template(template).render(**data)

    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"
    if DISCLAIMER.rstrip(".") not in text:
        text = f"> **{DISCLAIMER}**\n\n{text}"
    return text


# ---------------------------------------------------------------------------
# Markdown → HTML (minimal, dependency-free; enough for the plan template)
# ---------------------------------------------------------------------------

def _inline(text: str) -> str:
    out = html.escape(text, quote=False)
    out = out.replace("&lt;br&gt;", "<br>")
    out = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", out)
    out = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<em>\1</em>", out)
    out = re.sub(r"`([^`]+)`", r"<code>\1</code>", out)
    return re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r'<a href="\2">\1</a>', out)


def _table_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table_start(lines: list[str], i: int) -> bool:
    return lines[i].strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[i + 1].strip()) is not None


def markdown_to_html(markdown: str, title: str = "Business plan") -> str:
    lines = markdown.splitlines()
    parts: list[str] = []
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            parts.append(f"<h{level}>{_inline(heading.group(2))}</h{level}>")
            i += 1
            continue
        if stripped in ("---", "***"):
            parts.append("<hr>")
            i += 1
            continue
        if _is_table_start(lines, i):
            header = _table_cells(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_table_cells(lines[i]))
                i += 1
            head = "".join(f"<th>{_inline(c)}</th>" for c in header)
            body = "".join("<tr>" + "".join(f"<td>{_inline(c)}</td>" for c in row) + "</tr>" for row in rows)
            parts.append(f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
            continue
        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            parts.append(f"<blockquote>{_inline(' '.join(quote))}</blockquote>")
            continue
        if re.match(r"^[-*]\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]))
                i += 1
            parts.append("<ul>" + "".join(f"<li>{_inline(item)}</li>" for item in items) + "</ul>")
            continue
        paragraph = []
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6}\s|\||>|[-*]\s|---$)", lines[i].strip()):
            paragraph.append(lines[i].strip())
            i += 1
        if not paragraph:  # a lone line starting with a block marker we don't render specially
            paragraph.append(stripped)
            i += 1
        parts.append(f"<p>{_inline(' '.join(paragraph))}</p>")
    style = (
        "body{font-family:Georgia,serif;max-width:46rem;margin:2rem auto;line-height:1.5;color:#1b1b1b}"
        "table{border-collapse:collapse;width:100%;margin:1rem 0}"
        "td,th{border:1px solid #ccc;padding:.35rem .5rem;text-align:left;vertical-align:top}"
        "blockquote{border-left:4px solid #999;margin:1rem 0;padding:.25rem 1rem;background:#f6f6f6}"
    )
    return (
        '<!doctype html><html><head><meta charset="utf-8">'
        f"<title>{html.escape(title)}</title><style>{style}</style></head>"
        f"<body>{''.join(parts)}</body></html>"
    )


# ---------------------------------------------------------------------------
# Optional renderers
# ---------------------------------------------------------------------------

def _docx_importable() -> bool:
    return importlib.util.find_spec("docx") is not None


def _weasyprint_importable() -> bool:
    return importlib.util.find_spec("weasyprint") is not None


def capabilities() -> dict[str, Any]:
    pandoc = shutil.which("pandoc")
    pdf_engine = "weasyprint" if _weasyprint_importable() else ("pandoc" if pandoc else None)
    return {"md": True, "docx": _docx_importable(), "pdf": pdf_engine is not None, "pdf_engine": pdf_engine}


def _runs(text: str) -> list[tuple[str, bool, bool]]:
    """Split ``**bold**`` / ``*italic*`` into (text, bold, italic) runs."""
    runs: list[tuple[str, bool, bool]] = []
    for part in re.split(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)", text.replace("<br>", "\n")):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append((part[1:-1], False, True))
        else:
            runs.append((part, False, False))
    return runs


def markdown_to_docx(markdown: str) -> tuple[bytes | None, str | None]:
    if not _docx_importable():
        return None, "python-docx is not installed"
    import io

    import docx  # type: ignore

    document = docx.Document()

    def add_runs(paragraph: Any, text: str) -> None:
        for chunk, bold, italic in _runs(text):
            run = paragraph.add_run(chunk)
            run.bold = bold
            run.italic = italic

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped or stripped in ("---", "***"):
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
            i += 1
            continue
        if _is_table_start(lines, i):
            rows = [_table_cells(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_table_cells(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            table = document.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r_index, row in enumerate(rows):
                for c_index in range(width):
                    add_runs(table.cell(r_index, c_index).paragraphs[0], row[c_index] if c_index < len(row) else "")
            continue
        if re.match(r"^[-*]\s+", stripped):
            add_runs(document.add_paragraph(style="List Bullet"), re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue
        add_runs(document.add_paragraph(), stripped.lstrip(">").strip() if stripped.startswith(">") else stripped)
        i += 1

    document.core_properties.comments = DISCLAIMER
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), None


def markdown_to_pdf(markdown: str, title: str = "Business plan") -> tuple[bytes | None, str | None]:
    if _weasyprint_importable():
        try:
            from weasyprint import HTML  # type: ignore

            return HTML(string=markdown_to_html(markdown, title)).write_pdf(), None
        except Exception as exc:  # pragma: no cover - depends on system libraries
            return None, f"weasyprint failed: {exc}"
    pandoc = shutil.which("pandoc")
    if pandoc:
        with tempfile.TemporaryDirectory() as tmp:
            source, target = Path(tmp) / "plan.md", Path(tmp) / "plan.pdf"
            source.write_text(markdown, encoding="utf-8")
            try:
                subprocess.run([pandoc, str(source), "-o", str(target)], check=True, capture_output=True, timeout=120)
                return target.read_bytes(), None
            except (subprocess.SubprocessError, OSError) as exc:  # pragma: no cover - needs a PDF engine
                return None, f"pandoc failed: {exc}"
    return None, "no PDF renderer available (install weasyprint, or pandoc with a PDF engine)"


def render(
    context: Mapping[str, Any],
    formats: Sequence[str] = ("md",),
    template: str = DEFAULT_TEMPLATE,
    *,
    template_dir: str | Path | None = None,
) -> dict[str, Any]:
    wanted = [str(f).lower() for f in (formats or ["md"])]
    unknown = sorted(set(wanted) - set(SUPPORTED_FORMATS))
    if unknown:
        raise DocgenError(f"Unsupported formats: {', '.join(unknown)}")

    markdown = render_markdown(context, template, template_dir=template_dir)
    business = context.get("business") if isinstance(context.get("business"), Mapping) else {}
    title = str((business or {}).get("name") or "Business plan")
    result: dict[str, Any] = {
        "template": template,
        "markdown": markdown,
        "formats": {"md": {"ok": True, "reason": None}},
        "disclaimer": DISCLAIMER,
    }
    if "docx" in wanted:
        data, reason = markdown_to_docx(markdown)
        result["formats"]["docx"] = {"ok": data is not None, "reason": reason}
        if data is not None:
            result["docx_b64"] = base64.b64encode(data).decode("ascii")
    if "pdf" in wanted:
        data, reason = markdown_to_pdf(markdown, title)
        result["formats"]["pdf"] = {"ok": data is not None, "reason": reason}
        if data is not None:
            result["pdf_b64"] = base64.b64encode(data).decode("ascii")
    return result
