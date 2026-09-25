"""Deterministic finance model: assumptions → workbook with live formulas + summary.

No LLM is involved anywhere in this module. Given ``finance/assumptions.yaml``
(the founder's own numbers), it produces:

1. ``project()`` — a month-by-month projection computed in plain Python;
2. ``summarize()`` — the headline figures the business plan is allowed to
   quote: revenue, gross margin, EBITDA, cash low point, runway months and the
   break-even month;
3. ``build_workbook_spec()`` — the workbook layout from
   ``templates/finance-model.yaml`` with LIVE formulas (sheets Assumptions,
   Revenue, Costs, P&L, CashFlow, Balance);
4. ``verify_workbook()`` — evaluates those formulas with ``formula_eval`` and
   checks every monthly cell against the projection, so the spreadsheet and
   the summary cannot drift;
5. ``workbook_bytes()`` — writes the .xlsx with openpyxl, imported lazily; when
   openpyxl is missing the model still returns the summary and says why the
   workbook is unavailable.

Model shape: pre-tax, cash basis, no debtors, creditors or depreciation;
set-up costs are expensed in month 1. See the template header for formulas.

Not legal or financial advice.
"""

from __future__ import annotations

import base64
import hashlib
import io
import math
import re
from dataclasses import asdict, dataclass
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, Mapping

import yaml

from . import DISCLAIMER, pack_dir
from .formula_eval import Workbook, column_index, column_letter

__all__ = [
    "AssumptionError",
    "Assumptions",
    "MonthRow",
    "build_workbook_spec",
    "load_template",
    "project",
    "round2",
    "run_model",
    "summarize",
    "verify_workbook",
    "workbook_bytes",
]

TEMPLATE_FILE = "templates/finance-model.yaml"
MODEL_VERSION = "1.0.0"
_EPS = 1e-9
_CURRENCY_RE = re.compile(r"^[A-Z]{3}$")


class AssumptionError(ValueError):
    """Assumptions are missing, non-numeric or out of range."""

    def __init__(self, errors: dict[str, str]):
        self.errors = errors
        super().__init__("; ".join(f"{key}: {message}" for key, message in errors.items()))


def round2(value: float) -> float:
    """Half-up rounding to cents (not banker's rounding)."""
    return float(Decimal(repr(float(value))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


@dataclass(frozen=True)
class Assumptions:
    currency: str
    starting_cash: float
    setup_costs: float
    price_per_unit: float
    units_month_one: float
    monthly_growth_pct: float
    cost_of_sale_pct: float
    fixed_monthly_costs: float
    marketing_monthly: float = 0.0
    months: int = 36

    # (key, minimum, maximum, required)
    _NUMERIC = (
        ("starting_cash", 0.0, 1e12, True),
        ("setup_costs", 0.0, 1e12, True),
        ("price_per_unit", 0.0, 1e9, True),
        ("units_month_one", 0.0, 1e9, True),
        ("monthly_growth_pct", -99.0, 1000.0, True),
        ("cost_of_sale_pct", 0.0, 500.0, True),
        ("fixed_monthly_costs", 0.0, 1e12, True),
        ("marketing_monthly", 0.0, 1e12, False),
    )

    @classmethod
    def from_mapping(cls, data: Mapping[str, Any], *, default_months: int = 36, max_months: int = 120) -> "Assumptions":
        errors: dict[str, str] = {}
        values: dict[str, Any] = {}

        currency = str(data.get("currency") or "").strip().upper()
        if not _CURRENCY_RE.match(currency):
            errors["currency"] = "must be a three-letter currency code such as GBP, ZAR or USD"
        values["currency"] = currency

        for key, minimum, maximum, required in cls._NUMERIC:
            raw = data.get(key)
            if raw is None or raw == "":
                if required:
                    errors[key] = "is required"
                values[key] = 0.0
                continue
            try:
                number = float(raw)
            except (TypeError, ValueError):
                errors[key] = "must be a number"
                continue
            if not math.isfinite(number):
                errors[key] = "must be a finite number"
            elif number < minimum or number > maximum:
                errors[key] = f"must be between {minimum:g} and {maximum:g}"
            values[key] = number

        raw_months = data.get("months")
        months = default_months
        if raw_months not in (None, ""):
            try:
                as_float = float(raw_months)
                if not as_float.is_integer():
                    raise ValueError
                months = int(as_float)
            except (TypeError, ValueError):
                errors["months"] = "must be a whole number"
        if "months" not in errors and not 1 <= months <= max_months:
            errors["months"] = f"must be between 1 and {max_months}"
        values["months"] = months

        if errors:
            raise AssumptionError(errors)
        return cls(**values)

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class MonthRow:
    month: int
    units: float
    price: float
    revenue: float
    cost_of_sales: float
    gross_profit: float
    fixed_costs: float
    marketing: float
    setup_costs: float
    operating_costs: float
    total_costs: float
    ebitda: float
    opening_cash: float
    closing_cash: float
    share_capital: float
    retained_earnings: float
    balance_check: float


def project(a: Assumptions) -> list[MonthRow]:
    """Month-by-month projection in plain Python (the reference the workbook must match)."""
    growth = a.monthly_growth_pct / 100.0
    cost_share = a.cost_of_sale_pct / 100.0
    rows: list[MonthRow] = []
    cash = a.starting_cash
    retained = 0.0
    for month in range(1, a.months + 1):
        units = a.units_month_one * math.pow(1.0 + growth, month - 1)
        revenue = units * a.price_per_unit
        cost_of_sales = revenue * cost_share
        gross = revenue - cost_of_sales
        setup = a.setup_costs if month == 1 else 0.0
        operating = a.fixed_monthly_costs + a.marketing_monthly + setup
        ebitda = gross - operating
        opening = cash
        cash = opening + ebitda
        retained += ebitda
        rows.append(
            MonthRow(
                month=month,
                units=units,
                price=a.price_per_unit,
                revenue=revenue,
                cost_of_sales=cost_of_sales,
                gross_profit=gross,
                fixed_costs=a.fixed_monthly_costs,
                marketing=a.marketing_monthly,
                setup_costs=setup,
                operating_costs=operating,
                total_costs=cost_of_sales + operating,
                ebitda=ebitda,
                opening_cash=opening,
                closing_cash=cash,
                share_capital=a.starting_cash,
                retained_earnings=retained,
                balance_check=cash - (a.starting_cash + retained),
            )
        )
    return rows


def _margin_pct(gross: float, revenue: float) -> float | None:
    return round2(100.0 * gross / revenue) if revenue > _EPS else None


def summarize(a: Assumptions, rows: list[MonthRow]) -> dict[str, Any]:
    """Headline figures — the only numbers the business plan may quote."""
    if not rows:
        raise ValueError("No months to summarise")

    year1 = rows[:12]
    y1_revenue = math.fsum(r.revenue for r in year1)
    y1_gross = math.fsum(r.gross_profit for r in year1)
    y1_ebitda = math.fsum(r.ebitda for r in year1)

    low = min(rows, key=lambda r: (r.closing_cash, r.month))
    runs_out_row = next((r for r in rows if r.closing_cash < -_EPS), None)
    breakeven = next((r for r in rows if r.ebitda >= -_EPS), None)

    years = []
    for start in range(0, len(rows), 12):
        chunk = rows[start:start + 12]
        revenue = math.fsum(r.revenue for r in chunk)
        gross = math.fsum(r.gross_profit for r in chunk)
        years.append({
            "year": start // 12 + 1,
            "months": len(chunk),
            "revenue": round2(revenue),
            "gross_profit": round2(gross),
            "gross_margin_pct": _margin_pct(gross, revenue),
            "ebitda": round2(math.fsum(r.ebitda for r in chunk)),
            "closing_cash": round2(chunk[-1].closing_cash),
        })

    return {
        "model_version": MODEL_VERSION,
        "currency": a.currency,
        "months": a.months,
        "year1_revenue": round2(y1_revenue),
        "year1_gross_profit": round2(y1_gross),
        "year1_gross_margin_pct": _margin_pct(y1_gross, y1_revenue),
        "year1_ebitda": round2(y1_ebitda),
        "total_revenue": round2(math.fsum(r.revenue for r in rows)),
        "total_ebitda": round2(math.fsum(r.ebitda for r in rows)),
        "month1_burn": round2(max(0.0, -rows[0].ebitda)),
        "cash_low_point": round2(low.closing_cash),
        "cash_low_month": low.month,
        "runs_out": runs_out_row is not None,
        "runway_months": (runs_out_row.month - 1) if runs_out_row is not None else a.months,
        "breakeven_month": breakeven.month if breakeven is not None else None,
        "closing_cash": round2(rows[-1].closing_cash),
        "years": years,
        "disclaimer": DISCLAIMER,
    }


# ---------------------------------------------------------------------------
# Workbook layout
# ---------------------------------------------------------------------------

def load_template(path: str | Path | None = None) -> dict[str, Any]:
    template_path = Path(path) if path else pack_dir() / TEMPLATE_FILE
    with open(template_path, encoding="utf-8") as handle:
        data = yaml.safe_load(handle)
    if not isinstance(data, dict) or "sheets" not in data or "assumptions" not in data:
        raise ValueError(f"Not a finance-model template: {template_path}")
    return data


def _sheet_ref(name: str) -> str:
    return name if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name) else "'" + name.replace("'", "''") + "'"


def build_workbook_spec(a: Assumptions, template: Mapping[str, Any] | None = None) -> dict[str, Any]:
    """Sheets as ``{name: [[cell, ...], ...]}`` plus per-cell formats.

    Cells are values or formula strings (``=...``) exactly as they will be
    written to the .xlsx.
    """
    template = template or load_template()
    sheets: dict[str, list[list[Any]]] = {}
    formats: dict[str, dict[str, str]] = {}

    # Assumptions sheet --------------------------------------------------
    spec = template["assumptions"]
    a_name = spec.get("sheet", "Assumptions")
    first = int(spec.get("first_row", 2))
    rows: list[list[Any]] = [list(spec.get("header", ["Assumption", "Value", "Notes"]))]
    while len(rows) < first - 1:
        rows.append(["", "", ""])
    refs: dict[str, str] = {}
    a_formats: dict[str, str] = {}
    values = a.as_dict()
    for offset, row in enumerate(spec["rows"]):
        key = row["key"]
        if key not in values:
            raise ValueError(f"Template assumption {key!r} is not a model input")
        value: Any = values[key]
        fmt = row.get("format", "number")
        if fmt == "percent":
            value = value / 100.0
        excel_row = first + offset
        refs[key] = f"{_sheet_ref(a_name)}!$B${excel_row}"
        a_formats[f"B{excel_row}"] = fmt
        rows.append([row.get("label", key), value, row.get("note", "")])
    rows.append(["", "", ""])
    rows.append([template.get("disclaimer", DISCLAIMER), "", ""])
    sheets[a_name] = rows
    formats[a_name] = a_formats

    # Monthly sheets -----------------------------------------------------
    first_data, last_data = 2, a.months + 1

    def expand(formula: str, row: int, month: int) -> Any:
        text = str(formula)
        if text == "{month}":
            return month
        for key, ref in refs.items():
            text = text.replace("{A:" + key + "}", ref)
        if "{A:" in text:
            raise ValueError(f"Unknown assumption placeholder in {formula!r}")
        return (
            text.replace("{r}", str(row))
            .replace("{prev}", str(row - 1))
            .replace("{first}", str(first_data))
            .replace("{last}", str(last_data))
        )

    for sheet in template["sheets"]:
        name = sheet["name"]
        columns = sheet["columns"]
        grid: list[list[Any]] = [[col["header"] for col in columns]]
        sheet_formats: dict[str, str] = {}
        for month in range(1, a.months + 1):
            row = month + 1
            cells = []
            for index, col in enumerate(columns, start=1):
                formula = col.get("first") if (month == 1 and col.get("first")) else col["formula"]
                cells.append(expand(formula, row, month))
                sheet_formats[f"{column_letter(index)}{row}"] = col.get("format", "number")
            grid.append(cells)

        totals = sheet.get("totals")
        if totals:
            row = last_data + 1
            cells = [""] * len(columns)
            cells[0] = totals.get("label", "Total")
            for letter in totals.get("columns", []):
                index = column_index(letter)
                cells[index - 1] = f"=SUM({letter}{first_data}:{letter}{last_data})"
                sheet_formats[f"{letter}{row}"] = columns[index - 1].get("format", "number")
            grid.append(cells)

        for footer in sheet.get("footer", []) or []:
            row = len(grid) + 2
            while len(grid) < row - 1:
                grid.append([""] * len(columns))
            cells = [""] * len(columns)
            cells[0] = footer["label"]
            cells[1] = expand(footer["formula"], row, 0)
            sheet_formats[f"B{row}"] = footer.get("format", "number")
            grid.append(cells)

        sheets[name] = grid
        formats[name] = sheet_formats

    return {
        "sheets": sheets,
        "formats": formats,
        "order": [a_name] + [s["name"] for s in template["sheets"]],
        "data_rows": {"first": first_data, "last": last_data},
        "title": template.get("title", "Finance model"),
    }


# Column in each monthly sheet → MonthRow attribute it must equal.
_CHECKS: dict[str, dict[str, str]] = {
    "Revenue": {"B": "units", "C": "price", "D": "revenue"},
    "Costs": {"B": "cost_of_sales", "C": "fixed_costs", "D": "marketing", "E": "setup_costs", "F": "total_costs"},
    "P&L": {"B": "revenue", "C": "cost_of_sales", "D": "gross_profit", "E": "operating_costs", "F": "ebitda"},
    "CashFlow": {"B": "opening_cash", "C": "ebitda", "D": "closing_cash"},
    "Balance": {"B": "closing_cash", "D": "share_capital", "E": "retained_earnings", "G": "balance_check"},
}


def verify_workbook(spec: Mapping[str, Any], rows: list[MonthRow], *, tolerance: float = 1e-6) -> dict[str, Any]:
    """Evaluate the workbook's formulas and compare every checked cell to the projection."""
    book = Workbook(spec["sheets"])
    mismatches: list[dict[str, Any]] = []
    checked = 0
    for sheet, columns in _CHECKS.items():
        if sheet not in spec["sheets"]:
            mismatches.append({"sheet": sheet, "cell": None, "reason": "missing sheet"})
            continue
        for row in rows:
            excel_row = row.month + 1
            for letter, attribute in columns.items():
                expected = getattr(row, attribute)
                actual = book.value(sheet, f"{letter}{excel_row}")
                checked += 1
                scale = max(1.0, abs(expected))
                if not isinstance(actual, (int, float)) or abs(float(actual) - expected) > tolerance * scale:
                    mismatches.append({"sheet": sheet, "cell": f"{letter}{excel_row}", "expected": expected, "actual": actual})
    low_row = len(spec["sheets"].get("CashFlow", [])) if "CashFlow" in spec["sheets"] else 0
    low_value = book.value("CashFlow", f"B{low_row}") if low_row else None
    expected_low = min(r.closing_cash for r in rows)
    if low_value is None or abs(float(low_value) - expected_low) > tolerance * max(1.0, abs(expected_low)):
        mismatches.append({"sheet": "CashFlow", "cell": f"B{low_row}", "expected": expected_low, "actual": low_value})
    return {"verified": not mismatches, "cells_checked": checked + 1, "mismatches": mismatches[:20]}


_NUMBER_FORMATS = {
    "money": "#,##0.00",
    "number": "#,##0.00",
    "integer": "0",
    "percent": "0.00%",
    "text": "@",
}


def workbook_bytes(spec: Mapping[str, Any]) -> tuple[bytes | None, str | None]:
    """Write the .xlsx with openpyxl (lazy import). Returns ``(bytes, None)`` or ``(None, reason)``."""
    try:
        from openpyxl import Workbook as XlsxWorkbook  # type: ignore
        from openpyxl.styles import Font  # type: ignore
    except ImportError:
        return None, "openpyxl is not installed — the summary is still exact; install openpyxl to produce model.xlsx"

    book = XlsxWorkbook()
    book.remove(book.active)
    for name in spec["order"]:
        sheet = book.create_sheet(title=name)
        for r_index, cells in enumerate(spec["sheets"][name], start=1):
            for c_index, value in enumerate(cells, start=1):
                if value == "" or value is None:
                    continue
                cell = sheet.cell(row=r_index, column=c_index, value=value)
                fmt = spec["formats"].get(name, {}).get(f"{column_letter(c_index)}{r_index}")
                if fmt in _NUMBER_FORMATS:
                    cell.number_format = _NUMBER_FORMATS[fmt]
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for c_index in range(1, (sheet.max_column or 1) + 1):
            sheet.column_dimensions[column_letter(c_index)].width = 18 if c_index > 1 else 28
    book.properties.title = spec.get("title", "Finance model")
    book.properties.description = DISCLAIMER

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue(), None


def run_model(
    assumptions: Mapping[str, Any],
    *,
    include_workbook: bool = True,
    template: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Validate → project → summarise → lay out → verify → (write xlsx)."""
    template = template or load_template()
    a = Assumptions.from_mapping(
        assumptions,
        default_months=int(template.get("default_months", 36)),
        max_months=int(template.get("max_months", 120)),
    )
    rows = project(a)
    summary = summarize(a, rows)
    spec = build_workbook_spec(a, template)
    verification = verify_workbook(spec, rows)

    workbook: dict[str, Any] = {"format": "xlsx", "available": False, "reason": "not requested"}
    if include_workbook:
        data, reason = workbook_bytes(spec)
        if data is None:
            workbook = {"format": "xlsx", "available": False, "reason": reason}
        else:
            workbook = {
                "format": "xlsx",
                "available": True,
                "bytes": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
                "b64": base64.b64encode(data).decode("ascii"),
            }

    return {
        "assumptions": a.as_dict(),
        "summary": summary,
        "monthly": [
            {key: (round2(value) if isinstance(value, float) else value) for key, value in asdict(row).items()}
            for row in rows
        ],
        "formulas_verified": verification["verified"],
        "verification": verification,
        "workbook": workbook,
        "disclaimer": DISCLAIMER,
    }
